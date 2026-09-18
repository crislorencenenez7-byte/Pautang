from flask import Flask, request, jsonify, send_from_directory
from openpyxl import load_workbook
from datetime import date, timedelta
from pathlib import Path
from threading import Lock

BASE = Path(__file__).resolve().parent
EXCEL_FILE = BASE / "PAUTANG.xlsx"
app = Flask(__name__, static_folder=str(BASE))
excel_lock = Lock()

@app.get("/")
def home():
    return send_from_directory(BASE, "index.html")

@app.post("/api/apply")
def apply_loan():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()

    try:
        amount = float(data.get("amount", 0))
    except (TypeError, ValueError):
        amount = 0

    if not name:
        return jsonify(error="Name is required."), 400
    if amount <= 0:
        return jsonify(error="Amount must be greater than 0."), 400

    loan_date = date.today()
    due_date = loan_date + timedelta(days=15)

    with excel_lock:
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active

        # Expected columns:
        # A Name | B Amount | C Date | D Due Date | E Status
        if ws.max_row == 1 and all(ws.cell(1, c).value is None for c in range(1, 6)):
            headers = ["Name", "Amount", "Date", "Due Date", "Status"]
            for c, header in enumerate(headers, 1):
                ws.cell(1, c).value = header

        # Ensure the expected header names exist.
        headers = ["Name", "Amount", "Date", "Due Date", "Status"]
        for c, header in enumerate(headers, 1):
            ws.cell(1, c).value = header

        row = ws.max_row + 1
        ws.cell(row, 1).value = name
        ws.cell(row, 2).value = amount
        ws.cell(row, 3).value = loan_date
        ws.cell(row, 4).value = due_date
        ws.cell(row, 5).value = "Unpaid"

        ws.cell(row, 2).number_format = '₱#,##0.00'
        ws.cell(row, 3).number_format = 'mm/dd/yyyy'
        ws.cell(row, 4).number_format = 'mm/dd/yyyy'

        wb.save(EXCEL_FILE)

    return jsonify(
        success=True,
        date=loan_date.strftime("%m/%d/%Y"),
        due_date=due_date.strftime("%m/%d/%Y"),
        status="Unpaid"
    )

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
